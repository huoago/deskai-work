from __future__ import annotations

import hashlib
import re
import uuid
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import Workbook

from app.database.models import GeneratedArtifact
from app.database.session import Database

SAFE_NAME_RE = re.compile(r"[^\w\-. ()\u4e00-\u9fff]+", re.UNICODE)
FORMULA_PREFIXES = ("=", "+", "-", "@")


class ArtifactService:
    def __init__(self, database: Database, data_dir: Path) -> None:
        self.database = database
        self.root = (data_dir / "generated").resolve()
        self.root.mkdir(parents=True, exist_ok=True)

    def create_word_document(
        self,
        *,
        task_id: str,
        workspace_id: str,
        filename: str,
        title: str,
        sections: list[dict[str, Any]],
    ) -> GeneratedArtifact:
        target = self._target(task_id, filename, ".docx")
        document = Document()
        if title.strip():
            document.add_heading(title.strip()[:500], level=0)

        for raw in sections[:30]:
            heading = str(raw.get("heading") or "").strip()
            body = str(raw.get("body") or "")
            if heading:
                document.add_heading(heading[:500], level=1)
            for paragraph in body[:30000].splitlines() or [""]:
                document.add_paragraph(paragraph)

        document.save(target)
        return self._record(
            task_id=task_id,
            workspace_id=workspace_id,
            kind="docx",
            target=target,
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    def create_spreadsheet(
        self,
        *,
        task_id: str,
        workspace_id: str,
        filename: str,
        sheets: list[dict[str, Any]],
    ) -> GeneratedArtifact:
        target = self._target(task_id, filename, ".xlsx")
        workbook = Workbook()
        default = workbook.active
        workbook.remove(default)

        used_names: set[str] = set()
        for index, raw in enumerate(sheets[:10], start=1):
            name = self._sheet_name(str(raw.get("name") or f"Sheet{index}"), used_names)
            used_names.add(name)
            sheet = workbook.create_sheet(title=name)

            headers = list(raw.get("headers") or [])[:50]
            rows = list(raw.get("rows") or [])[:2000]
            if headers:
                sheet.append([self._safe_cell(value) for value in headers])
            for row in rows:
                values = list(row)[:50] if isinstance(row, list) else [row]
                sheet.append([self._safe_cell(value) for value in values])

            if headers:
                sheet.freeze_panes = "A2"
                sheet.auto_filter.ref = sheet.dimensions

        if not workbook.sheetnames:
            workbook.create_sheet(title="Sheet1")
        workbook.save(target)
        return self._record(
            task_id=task_id,
            workspace_id=workspace_id,
            kind="xlsx",
            target=target,
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def _target(self, task_id: str, requested: str, suffix: str) -> Path:
        safe_task = re.sub(r"[^A-Za-z0-9-]", "", task_id)
        if not safe_task:
            raise ValueError("Invalid task id")
        task_root = (self.root / safe_task).resolve()
        if not task_root.is_relative_to(self.root):
            raise ValueError("Generated artifact path escaped the DeskAI sandbox")
        task_root.mkdir(parents=True, exist_ok=True)

        base = Path(str(requested or "")).name
        stem = Path(base).stem or "deskai-output"
        stem = SAFE_NAME_RE.sub("_", stem).strip(" ._")[:120] or "deskai-output"
        target = (task_root / f"{stem}{suffix}").resolve()
        if not target.is_relative_to(task_root):
            raise ValueError("Generated artifact path escaped the task sandbox")
        if target.exists():
            target = task_root / f"{stem}-{uuid.uuid4().hex[:8]}{suffix}"
        return target

    @staticmethod
    def _sheet_name(value: str, used: set[str]) -> str:
        base = re.sub(r"[:\\/?*\[\]]", "_", value).strip()[:31] or "Sheet"
        candidate = base
        index = 2
        while candidate in used:
            suffix = f"-{index}"
            candidate = f"{base[:31-len(suffix)]}{suffix}"
            index += 1
        return candidate

    @staticmethod
    def _safe_cell(value: Any) -> str | int | float | bool | None:
        if value is None or isinstance(value, (int, float, bool)):
            return value
        text = str(value)[:10000]
        if text.startswith(FORMULA_PREFIXES):
            return "'" + text
        return text

    def _record(
        self,
        *,
        task_id: str,
        workspace_id: str,
        kind: str,
        target: Path,
        mime_type: str,
    ) -> GeneratedArtifact:
        digest = hashlib.sha256(target.read_bytes()).hexdigest()
        artifact = GeneratedArtifact(
            task_id=task_id,
            workspace_id=workspace_id,
            kind=kind,
            filename=target.name,
            path=str(target),
            mime_type=mime_type,
            sha256=digest,
            size=target.stat().st_size,
        )
        try:
            with self.database.session() as session:
                session.add(artifact)
                session.flush()
                session.expunge(artifact)
            return artifact
        except Exception:
            target.unlink(missing_ok=True)
            raise
