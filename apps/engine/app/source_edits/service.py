from __future__ import annotations

import difflib
import hashlib
import os
import shutil
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple
from sqlalchemy import select

from app.database.models import (
    AuditLog,
    File,
    SourceFileEdit,
    Task,
    WorkspaceRoot,
)
from app.database.session import Database
from app.indexing.hashing import sha256_file
from app.indexing.scanner import scan_root
from app.security.paths import normalize_root, require_within

SUPPORTED_EDIT_EXTENSIONS = {".txt", ".md", ".docx", ".xlsx"}
MAX_EDIT_SOURCE_SIZE = 100 * 1024 * 1024
MAX_TEXT_EDIT_SIZE = 10 * 1024 * 1024
MAX_REPLACEMENTS = 20
MAX_CELL_EDITS = 100
MAX_DIFF_PREVIEW = 20000
FORMULA_PREFIXES = ("=", "+", "-", "@")


class SourceFileEditService:
    def __init__(self, database: Database, data_dir: Path) -> None:
        self.database = database
        self.proposal_root = (data_dir / "edit_proposals").resolve()
        self.backup_root = (data_dir / "edit_backups").resolve()
        self.proposal_root.mkdir(parents=True, exist_ok=True)
        self.backup_root.mkdir(parents=True, exist_ok=True)

    def propose(
        self,
        *,
        task_id: str,
        workspace_id: str,
        file_id: str,
        mode: str,
        summary: str,
        replacements: list[dict[str, Any]],
        cell_edits: list[dict[str, Any]],
    ) -> dict[str, Any]:
        file, root, source = self._editable_file(
            task_id=task_id,
            workspace_id=workspace_id,
            file_id=file_id,
        )
        extension = (file.extension or source.suffix).lower()
        if extension not in SUPPORTED_EDIT_EXTENSIONS:
            raise ValueError("Phase 11 supports source edits only for TXT, MD, DOCX, and XLSX")
        if source.stat().st_size > MAX_EDIT_SOURCE_SIZE:
            raise ValueError("Source file exceeds the 100 MB edit limit")

        current_sha = sha256_file(source)
        if not file.sha256 or current_sha != file.sha256:
            raise ValueError("Source file changed outside DeskAI; rescan it before proposing an edit")

        normalized_summary = str(summary or "").strip()[:1000] or f"Proposed edit to {file.filename}"
        edit_id = str(uuid.uuid4())
        candidate_dir = (self.proposal_root / edit_id).resolve()
        if not candidate_dir.is_relative_to(self.proposal_root):
            raise ValueError("Edit proposal path escaped the DeskAI sandbox")
        candidate_dir.mkdir(parents=True, exist_ok=False)
        candidate = (candidate_dir / source.name).resolve()
        if not candidate.is_relative_to(candidate_dir):
            raise ValueError("Edit candidate path escaped the proposal sandbox")

        try:
            if mode == "text_replace":
                if extension not in {".txt", ".md"}:
                    raise ValueError("text_replace is limited to TXT/MD files")
                diff_preview, change_spec = self._propose_text(source, candidate, replacements)
                kind = "text"
            elif mode == "docx_replace":
                if extension != ".docx":
                    raise ValueError("docx_replace requires a DOCX file")
                diff_preview, change_spec = self._propose_docx(source, candidate, replacements)
                kind = "docx"
            elif mode == "xlsx_cells":
                if extension != ".xlsx":
                    raise ValueError("xlsx_cells requires an XLSX file")
                diff_preview, change_spec = self._propose_xlsx(source, candidate, cell_edits)
                kind = "xlsx"
            else:
                raise ValueError("Unsupported source-edit mode")

            candidate_sha = sha256_file(candidate)
            if candidate_sha == current_sha:
                raise ValueError("Proposed edit produced no file changes")

            edit = SourceFileEdit(
                id=edit_id,
                task_id=task_id,
                workspace_id=workspace_id,
                file_id=file_id,
                kind=kind,
                status="pending",
                summary=normalized_summary,
                diff_preview=diff_preview[:MAX_DIFF_PREVIEW],
                change_spec_json=change_spec,
                original_sha256=current_sha,
                candidate_sha256=candidate_sha,
                candidate_path=str(candidate),
            )
            with self.database.session() as session:
                session.add(edit)
                session.flush()
                session.expunge(edit)
            return self.payload(edit, filename=file.filename)
        except Exception:
            shutil.rmtree(candidate_dir, ignore_errors=True)
            raise

    def confirm(self, edit_id: str) -> dict[str, Any]:
        snapshot = self._action_snapshot(edit_id, expected_status="pending")
        source = snapshot["source"]
        candidate = snapshot["candidate"]
        backup = snapshot["backup"]
        root = snapshot["root"]

        current_sha = sha256_file(source)
        if current_sha != snapshot["original_sha256"]:
            raise ValueError("Source file changed after the proposal; confirmation is blocked")
        if not candidate.is_file() or sha256_file(candidate) != snapshot["candidate_sha256"]:
            raise ValueError("Staged edit candidate is missing or corrupted")

        backup.parent.mkdir(parents=True, exist_ok=True)
        if backup.exists():
            raise ValueError("Backup already exists for this edit proposal")
        shutil.copy2(source, backup)
        self._atomic_replace(candidate, source)

        applied_sha = sha256_file(source)
        if applied_sha != snapshot["candidate_sha256"]:
            self._atomic_replace(backup, source)
            raise ValueError("Applied file hash did not match the staged candidate")

        now = datetime.now(timezone.utc)
        scan_error: str | None = None
        try:
            with self.database.session() as session:
                edit = session.get(SourceFileEdit, edit_id)
                if edit is None or edit.status != "pending":
                    raise ValueError("Edit proposal is no longer pending")
                edit.status = "applied"
                edit.confirmed_at = now
                edit.applied_at = now
                edit.applied_sha256 = applied_sha
                edit.backup_path = str(backup)
                edit.error_message = None
                try:
                    scan_root(session, root, queue_changes=True, metadata_shortcut=False)
                except Exception as exc:
                    scan_error = f"{type(exc).__name__}: {exc}"[:2000]
                    edit.error_message = f"Applied successfully; reindex rescan failed: {scan_error}"
                session.add(
                    AuditLog(
                        task_id=edit.task_id,
                        action="source_edit_applied",
                        target=edit.file_id,
                        result=(
                            f"edit_id={edit.id}; backup_created=true"
                            + (f"; scan_error={scan_error}" if scan_error else "; reindex_scan=ok")
                        ),
                        risk_level=5,
                    )
                )
        except Exception:
            try:
                self._atomic_replace(backup, source)
            finally:
                raise

        return self.get(edit_id)

    def reject(self, edit_id: str) -> dict[str, Any]:
        now = datetime.now(timezone.utc)
        with self.database.session() as session:
            edit = session.get(SourceFileEdit, edit_id)
            if edit is None:
                raise ValueError("Edit proposal not found")
            if edit.status != "pending":
                raise ValueError("Only pending edit proposals can be rejected")
            edit.status = "rejected"
            edit.rejected_at = now
            session.add(
                AuditLog(
                    task_id=edit.task_id,
                    action="source_edit_rejected",
                    target=edit.file_id,
                    result=f"edit_id={edit.id}",
                    risk_level=3,
                )
            )
        return self.get(edit_id)

    def rollback(self, edit_id: str) -> dict[str, Any]:
        snapshot = self._action_snapshot(edit_id, expected_status="applied")
        source = snapshot["source"]
        candidate = snapshot["candidate"]
        backup = snapshot["backup"]
        root = snapshot["root"]

        if not snapshot["applied_sha256"]:
            raise ValueError("Applied edit hash is missing")
        if sha256_file(source) != snapshot["applied_sha256"]:
            raise ValueError("Source file changed after the edit; automatic rollback is blocked")
        if not backup.is_file() or sha256_file(backup) != snapshot["original_sha256"]:
            raise ValueError("Original backup is missing or corrupted")
        if not candidate.is_file() or sha256_file(candidate) != snapshot["candidate_sha256"]:
            raise ValueError("Applied candidate copy is missing or corrupted")

        self._atomic_replace(backup, source)
        restored_sha = sha256_file(source)
        if restored_sha != snapshot["original_sha256"]:
            self._atomic_replace(candidate, source)
            raise ValueError("Rollback hash verification failed")

        now = datetime.now(timezone.utc)
        try:
            with self.database.session() as session:
                edit = session.get(SourceFileEdit, edit_id)
                if edit is None or edit.status != "applied":
                    raise ValueError("Edit proposal is no longer applied")
                edit.status = "rolled_back"
                edit.rolled_back_at = now
                edit.error_message = None
                try:
                    scan_root(session, root, queue_changes=True, metadata_shortcut=False)
                except Exception as exc:
                    edit.error_message = (
                        f"Rollback succeeded; reindex rescan failed: {type(exc).__name__}: {exc}"
                    )[:2000]
                session.add(
                    AuditLog(
                        task_id=edit.task_id,
                        action="source_edit_rolled_back",
                        target=edit.file_id,
                        result=f"edit_id={edit.id}; restored_sha256={restored_sha}",
                        risk_level=5,
                    )
                )
        except Exception:
            try:
                self._atomic_replace(candidate, source)
            finally:
                raise

        return self.get(edit_id)

    def get(self, edit_id: str) -> dict[str, Any]:
        with self.database.session() as session:
            edit = session.get(SourceFileEdit, edit_id)
            if edit is None:
                raise ValueError("Edit proposal not found")
            file = session.get(File, edit.file_id)
            filename = file.filename if file is not None else "Unknown file"
            return self.payload(edit, filename=filename)

    def list(
        self,
        *,
        workspace_id: str | None = None,
        task_id: str | None = None,
        limit: int = 100,
    ) -> list[dict[str, Any]]:
        with self.database.session() as session:
            statement = (
                select(SourceFileEdit)
                .order_by(SourceFileEdit.created_at.desc())
                .limit(max(1, min(int(limit), 500)))
            )
            if workspace_id:
                statement = statement.where(SourceFileEdit.workspace_id == workspace_id)
            if task_id:
                statement = statement.where(SourceFileEdit.task_id == task_id)
            edits = list(session.scalars(statement).all())
            file_ids = {item.file_id for item in edits}
            names = {
                item.id: item.filename
                for item in session.scalars(select(File).where(File.id.in_(file_ids))).all()
            } if file_ids else {}
            return [self.payload(item, filename=names.get(item.file_id, "Unknown file")) for item in edits]

    @staticmethod
    def payload(edit: SourceFileEdit, *, filename: str) -> dict[str, Any]:
        return {
            "id": edit.id,
            "task_id": edit.task_id,
            "workspace_id": edit.workspace_id,
            "file_id": edit.file_id,
            "filename": filename,
            "kind": edit.kind,
            "status": edit.status,
            "summary": edit.summary,
            "diff_preview": edit.diff_preview,
            "original_sha256": edit.original_sha256,
            "candidate_sha256": edit.candidate_sha256,
            "applied_sha256": edit.applied_sha256,
            "error_message": edit.error_message,
            "created_at": edit.created_at.isoformat(),
            "confirmed_at": edit.confirmed_at.isoformat() if edit.confirmed_at else None,
            "applied_at": edit.applied_at.isoformat() if edit.applied_at else None,
            "rejected_at": edit.rejected_at.isoformat() if edit.rejected_at else None,
            "rolled_back_at": edit.rolled_back_at.isoformat() if edit.rolled_back_at else None,
            "requires_user_confirmation": edit.status == "pending",
            "can_rollback": edit.status == "applied",
            "backup_created": bool(edit.backup_path),
        }

    def _editable_file(
        self,
        *,
        task_id: str,
        workspace_id: str,
        file_id: str,
    ) -> tuple[File, WorkspaceRoot, Path]:
        with self.database.session() as session:
            task = session.get(Task, task_id)
            if task is None or task.workspace_id != workspace_id:
                raise ValueError("Task is not part of the active Workspace")
            file = session.get(File, file_id)
            if file is None or file.workspace_id != workspace_id:
                raise ValueError("File is not part of the active Workspace")
            if file.status in {"deleted", "revoked", "unsupported"}:
                raise ValueError("File is not currently editable")
            source = Path(file.path)
            if source.is_symlink():
                raise ValueError("Symlink source files cannot be edited")
            source = source.resolve(strict=True)
            roots = list(
                session.scalars(
                    select(WorkspaceRoot).where(
                        WorkspaceRoot.workspace_id == workspace_id,
                        WorkspaceRoot.read_allowed.is_(True),
                        WorkspaceRoot.write_allowed.is_(True),
                    )
                ).all()
            )
            matched: WorkspaceRoot | None = None
            for root in roots:
                try:
                    normalized = normalize_root(root.path)
                    require_within(source, normalized)
                    matched = root
                    break
                except (OSError, ValueError):
                    continue
            if matched is None:
                raise ValueError(
                    "Source editing requires write_allowed on the containing Workspace root"
                )
            session.expunge(file)
            session.expunge(matched)
            return file, matched, source

    def _action_snapshot(self, edit_id: str, *, expected_status: str) -> dict[str, Any]:
        with self.database.session() as session:
            edit = session.get(SourceFileEdit, edit_id)
            if edit is None:
                raise ValueError("Edit proposal not found")
            if edit.status != expected_status:
                raise ValueError(f"Edit proposal must be {expected_status}")
            file = session.get(File, edit.file_id)
            if file is None or file.workspace_id != edit.workspace_id:
                raise ValueError("Source file is no longer available in the Workspace")
            task = session.get(Task, edit.task_id)
            if task is None or task.workspace_id != edit.workspace_id:
                raise ValueError("Edit task is no longer available")
            source = Path(file.path)
            if source.is_symlink():
                raise ValueError("Symlink source files cannot be edited")
            source = source.resolve(strict=True)
            roots = list(
                session.scalars(
                    select(WorkspaceRoot).where(
                        WorkspaceRoot.workspace_id == edit.workspace_id,
                        WorkspaceRoot.read_allowed.is_(True),
                        WorkspaceRoot.write_allowed.is_(True),
                    )
                ).all()
            )
            root = None
            for item in roots:
                try:
                    require_within(source, normalize_root(item.path))
                    root = item
                    break
                except (OSError, ValueError):
                    continue
            if root is None:
                raise ValueError("Workspace write permission was removed before this action")
            backup = (self.backup_root / edit.id / source.name).resolve()
            backup_dir = backup.parent
            if not backup_dir.is_relative_to(self.backup_root):
                raise ValueError("Edit backup path escaped the DeskAI sandbox")
            candidate = Path(edit.candidate_path).resolve(strict=True)
            if not candidate.is_relative_to(self.proposal_root):
                raise ValueError("Edit candidate path escaped the DeskAI sandbox")
            session.expunge(root)
            return {
                "source": source,
                "candidate": candidate,
                "backup": backup,
                "root": root,
                "original_sha256": edit.original_sha256,
                "candidate_sha256": edit.candidate_sha256,
                "applied_sha256": edit.applied_sha256,
            }

    def _propose_text(
        self,
        source: Path,
        candidate: Path,
        replacements: list[dict[str, Any]],
    ) -> tuple[str, dict[str, Any]]:
        if source.stat().st_size > MAX_TEXT_EDIT_SIZE:
            raise ValueError("TXT/MD source exceeds the 10 MB text-edit limit")
        before = source.read_text(encoding="utf-8")
        after, applied = self._apply_text_replacements(before, replacements)
        candidate.write_text(after, encoding="utf-8")
        diff = "\n".join(
            difflib.unified_diff(
                before.splitlines(),
                after.splitlines(),
                fromfile=source.name,
                tofile=f"{source.name} (proposed)",
                lineterm="",
            )
        )
        return diff[:MAX_DIFF_PREVIEW], {"mode": "text_replace", "replacements": applied}

    def _propose_docx(
        self,
        source: Path,
        candidate: Path,
        replacements: list[dict[str, Any]],
    ) -> tuple[str, dict[str, Any]]:
        specs = self._normalize_replacements(replacements)
        document = Document(source)
        paragraphs = list(self._docx_paragraphs(document))
        preview: list[str] = []
        applied: list[dict[str, Any]] = []
        for spec in specs:
            remaining = None if spec["replace_all"] else 1
            count = 0
            for paragraph in paragraphs:
                if remaining == 0:
                    break
                before = paragraph.text
                occurrences = before.count(spec["find"])
                if occurrences == 0:
                    continue
                allowed = occurrences if remaining is None else min(occurrences, remaining)
                after = before.replace(spec["find"], spec["replace"], allowed)
                if after != before:
                    paragraph.text = after
                    count += allowed
                    if len(preview) < 30:
                        preview.append(
                            f"- {before[:300]}\n+ {after[:300]}"
                        )
                    if remaining is not None:
                        remaining -= allowed
            if count == 0:
                raise ValueError(f"DOCX search text was not found: {spec['find'][:120]}")
            applied.append({**spec, "applied_count": count})
        document.save(candidate)
        diff = "\n".join(preview) or "DOCX content changed"
        diff += (
            "\n\nNote: changed DOCX paragraphs are rewritten by python-docx; "
            "inline run formatting inside those changed paragraphs may be simplified."
        )
        return diff[:MAX_DIFF_PREVIEW], {"mode": "docx_replace", "replacements": applied}

    def _propose_xlsx(
        self,
        source: Path,
        candidate: Path,
        cell_edits: list[dict[str, Any]],
    ) -> tuple[str, dict[str, Any]]:
        if not isinstance(cell_edits, list) or not cell_edits:
            raise ValueError("At least one spreadsheet cell edit is required")
        if len(cell_edits) > MAX_CELL_EDITS:
            raise ValueError("A proposal may modify at most 100 spreadsheet cells")
        workbook = load_workbook(source, data_only=False, keep_links=True)
        applied: list[dict[str, Any]] = []
        preview: list[str] = []
        for raw in cell_edits:
            if not isinstance(raw, dict):
                raise ValueError("Each cell edit must be an object")
            sheet_name = str(raw.get("sheet") or "").strip()
            cell_ref = str(raw.get("cell") or "").strip().upper()
            value_type = str(raw.get("value_type") or "").strip()
            raw_value = str(raw.get("value") or "")
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"Worksheet not found: {sheet_name}")
            try:
                row, column = coordinate_to_tuple(cell_ref)
            except ValueError as exc:
                raise ValueError(f"Invalid cell reference: {cell_ref}") from exc
            if row < 1 or column < 1:
                raise ValueError(f"Invalid cell reference: {cell_ref}")
            value = self._xlsx_value(value_type, raw_value)
            sheet = workbook[sheet_name]
            cell = sheet[cell_ref]
            old = cell.value
            if old == value:
                raise ValueError(f"Spreadsheet edit makes no change at {sheet_name}!{cell_ref}")
            cell.value = value
            applied.append(
                {
                    "sheet": sheet_name,
                    "cell": cell_ref,
                    "value_type": value_type,
                    "value": raw_value[:10000],
                    "old_value": self._display_value(old),
                    "new_value": self._display_value(value),
                }
            )
            preview.append(
                f"{sheet_name}!{cell_ref}: {self._display_value(old)} -> {self._display_value(value)}"
            )
        workbook.save(candidate)
        return "\n".join(preview)[:MAX_DIFF_PREVIEW], {
            "mode": "xlsx_cells",
            "cell_edits": applied,
        }

    @staticmethod
    def _atomic_replace(source_copy: Path, destination: Path) -> None:
        temp = destination.with_name(f".{destination.name}.deskai-{uuid.uuid4().hex[:8]}.tmp")
        try:
            shutil.copy2(source_copy, temp)
            os.replace(temp, destination)
        finally:
            temp.unlink(missing_ok=True)

    @staticmethod
    def _normalize_replacements(
        replacements: list[dict[str, Any]],
    ) -> list[dict[str, Any]]:
        if not isinstance(replacements, list) or not replacements:
            raise ValueError("At least one text replacement is required")
        if len(replacements) > MAX_REPLACEMENTS:
            raise ValueError("A proposal may contain at most 20 text replacements")
        normalized: list[dict[str, Any]] = []
        for raw in replacements:
            if not isinstance(raw, dict):
                raise ValueError("Each replacement must be an object")
            find = str(raw.get("find") or "")
            replace = str(raw.get("replace") or "")
            replace_all = bool(raw.get("replace_all"))
            if not find:
                raise ValueError("Replacement search text cannot be empty")
            if len(find) > 10000 or len(replace) > 20000:
                raise ValueError("Replacement text exceeds the allowed size")
            normalized.append(
                {
                    "find": find,
                    "replace": replace,
                    "replace_all": replace_all,
                }
            )
        return normalized

    def _apply_text_replacements(
        self,
        text: str,
        replacements: list[dict[str, Any]],
    ) -> tuple[str, list[dict[str, Any]]]:
        normalized = self._normalize_replacements(replacements)
        result = text
        applied: list[dict[str, Any]] = []
        for spec in normalized:
            count = result.count(spec["find"])
            if count == 0:
                raise ValueError(f"Search text was not found: {spec['find'][:120]}")
            applied_count = count if spec["replace_all"] else 1
            result = result.replace(
                spec["find"],
                spec["replace"],
                -1 if spec["replace_all"] else 1,
            )
            applied.append({**spec, "applied_count": applied_count})
        return result, applied

    @staticmethod
    def _docx_paragraphs(document) -> Iterable[Any]:
        for paragraph in document.paragraphs:
            yield paragraph
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        yield paragraph
                    for nested in cell.tables:
                        for nested_row in nested.rows:
                            for nested_cell in nested_row.cells:
                                for paragraph in nested_cell.paragraphs:
                                    yield paragraph

    @staticmethod
    def _xlsx_value(value_type: str, raw_value: str) -> str | float | bool | None:
        if value_type == "blank":
            return None
        if value_type == "number":
            try:
                return float(raw_value)
            except ValueError as exc:
                raise ValueError(f"Invalid numeric spreadsheet value: {raw_value[:120]}") from exc
        if value_type == "boolean":
            normalized = raw_value.strip().lower()
            if normalized not in {"true", "false"}:
                raise ValueError("Boolean spreadsheet values must be true or false")
            return normalized == "true"
        if value_type != "text":
            raise ValueError("Spreadsheet value_type must be text, number, boolean, or blank")
        text = raw_value[:10000]
        if text.startswith(FORMULA_PREFIXES):
            return "'" + text
        return text

    @staticmethod
    def _display_value(value: Any) -> str:
        if value is None:
            return "<blank>"
        text = str(value)
        return text if len(text) <= 300 else text[:297] + "..."
