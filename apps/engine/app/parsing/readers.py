from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Callable

import fitz
from docx import Document
from openpyxl import load_workbook
from PIL import ExifTags, Image
from pptx import Presentation

from app.parsing.schema import ParsedDocument, ParsedUnit

PARSER_VERSION = "phase3.1"


class ParseError(RuntimeError):
    pass


def _decode_text(path: Path) -> str:
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


def parse_text(path: Path) -> ParsedDocument:
    text = _decode_text(path).replace("\r\n", "\n").replace("\r", "\n")
    units: list[ParsedUnit] = []
    lines = text.split("\n")
    block: list[str] = []
    start = 1
    for index, line in enumerate(lines, start=1):
        if line.strip():
            if not block:
                start = index
            block.append(line)
        elif block:
            units.append(
                ParsedUnit(
                    kind="text_block",
                    text="\n".join(block),
                    locator={"line_start": start, "line_end": index - 1},
                )
            )
            block = []
    if block:
        units.append(
            ParsedUnit(
                kind="text_block",
                text="\n".join(block),
                locator={"line_start": start, "line_end": len(lines)},
            )
        )
    return ParsedDocument(
        parser="plain-text",
        parser_version=PARSER_VERSION,
        file_type=path.suffix.lower().lstrip("."),
        title=path.stem,
        text=text,
        units=units,
        metadata={"line_count": len(lines)},
    )


def parse_csv(path: Path) -> ParsedDocument:
    text = _decode_text(path)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows = list(reader)
    units: list[ParsedUnit] = []
    rendered: list[str] = []
    batch_size = 100
    for start in range(0, len(rows), batch_size):
        batch = rows[start : start + batch_size]
        block = "\n".join("\t".join(str(cell) for cell in row) for row in batch)
        rendered.append(block)
        units.append(
            ParsedUnit(
                kind="csv_rows",
                text=block,
                locator={"row_start": start + 1, "row_end": start + len(batch)},
            )
        )
    return ParsedDocument(
        parser="csv",
        parser_version=PARSER_VERSION,
        file_type="csv",
        title=path.stem,
        text="\n".join(rendered),
        units=units,
        metadata={"row_count": len(rows), "delimiter": getattr(dialect, "delimiter", ",")},
    )


def parse_pdf(path: Path) -> ParsedDocument:
    units: list[ParsedUnit] = []
    metadata: dict[str, object] = {}
    try:
        with fitz.open(path) as document:
            metadata = {key: value for key, value in (document.metadata or {}).items() if value}
            for page_index, page in enumerate(document, start=1):
                page_text = page.get_text("text", sort=True).strip()
                units.append(
                    ParsedUnit(
                        kind="page",
                        text=page_text,
                        locator={"page": page_index},
                        metadata={"width": float(page.rect.width), "height": float(page.rect.height)},
                    )
                )
    except Exception as exc:
        raise ParseError(f"PDF parse failed: {exc}") from exc
    text = "\n\n".join(unit.text for unit in units if unit.text)
    return ParsedDocument(
        parser="pymupdf",
        parser_version=PARSER_VERSION,
        file_type="pdf",
        title=str(metadata.get("title") or path.stem),
        text=text,
        units=units,
        metadata=metadata,
    )


def parse_docx(path: Path) -> ParsedDocument:
    try:
        document = Document(path)
    except Exception as exc:
        raise ParseError(f"DOCX parse failed: {exc}") from exc

    units: list[ParsedUnit] = []
    rendered: list[str] = []
    paragraph_number = 0
    for paragraph in document.paragraphs:
        paragraph_number += 1
        value = paragraph.text.strip()
        if not value:
            continue
        style_name = paragraph.style.name if paragraph.style is not None else None
        units.append(
            ParsedUnit(
                kind="paragraph",
                text=value,
                locator={"paragraph": paragraph_number},
                metadata={"style": style_name},
            )
        )
        rendered.append(value)

    for table_index, table in enumerate(document.tables, start=1):
        table_lines: list[str] = []
        for row_index, row in enumerate(table.rows, start=1):
            line = "\t".join(cell.text.strip() for cell in row.cells)
            table_lines.append(line)
            units.append(
                ParsedUnit(
                    kind="table_row",
                    text=line,
                    locator={"table": table_index, "row": row_index},
                )
            )
        rendered.append("\n".join(table_lines))

    properties = document.core_properties
    metadata = {
        "author": properties.author,
        "subject": properties.subject,
        "keywords": properties.keywords,
        "comments": properties.comments,
        "table_count": len(document.tables),
        "paragraph_count": len(document.paragraphs),
    }
    metadata = {key: value for key, value in metadata.items() if value not in (None, "")}
    return ParsedDocument(
        parser="python-docx",
        parser_version=PARSER_VERSION,
        file_type="docx",
        title=properties.title or path.stem,
        text="\n\n".join(part for part in rendered if part),
        units=units,
        metadata=metadata,
    )


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value)


def parse_xlsx(path: Path) -> ParsedDocument:
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:
        raise ParseError(f"XLSX parse failed: {exc}") from exc

    units: list[ParsedUnit] = []
    rendered_sheets: list[str] = []
    sheet_names: list[str] = []
    try:
        for sheet in workbook.worksheets:
            sheet_names.append(sheet.title)
            lines: list[str] = []
            first_nonempty: int | None = None
            last_nonempty: int | None = None
            for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = [_cell_text(value) for value in row]
                if not any(value for value in values):
                    continue
                line = "\t".join(values).rstrip()
                lines.append(line)
                if first_nonempty is None:
                    first_nonempty = row_index
                last_nonempty = row_index
            block = "\n".join(lines)
            if block:
                units.append(
                    ParsedUnit(
                        kind="sheet",
                        text=block,
                        locator={
                            "sheet": sheet.title,
                            "row_start": first_nonempty,
                            "row_end": last_nonempty,
                        },
                    )
                )
                rendered_sheets.append(f"[Sheet: {sheet.title}]\n{block}")
    finally:
        workbook.close()

    return ParsedDocument(
        parser="openpyxl",
        parser_version=PARSER_VERSION,
        file_type="xlsx",
        title=path.stem,
        text="\n\n".join(rendered_sheets),
        units=units,
        metadata={"sheet_names": sheet_names, "sheet_count": len(sheet_names)},
    )


def parse_pptx(path: Path) -> ParsedDocument:
    try:
        presentation = Presentation(path)
    except Exception as exc:
        raise ParseError(f"PPTX parse failed: {exc}") from exc

    units: list[ParsedUnit] = []
    rendered: list[str] = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                value = str(shape.text).strip()
                if value:
                    parts.append(value)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    line = "\t".join(cell.text.strip() for cell in row.cells)
                    if line.strip():
                        parts.append(line)
        slide_text = "\n".join(parts)
        units.append(
            ParsedUnit(
                kind="slide",
                text=slide_text,
                locator={"slide": slide_index},
            )
        )
        if slide_text:
            rendered.append(f"[Slide {slide_index}]\n{slide_text}")

    return ParsedDocument(
        parser="python-pptx",
        parser_version=PARSER_VERSION,
        file_type="pptx",
        title=path.stem,
        text="\n\n".join(rendered),
        units=units,
        metadata={"slide_count": len(presentation.slides)},
    )


def parse_image(path: Path) -> ParsedDocument:
    try:
        with Image.open(path) as image:
            exif_values: dict[str, str] = {}
            try:
                exif = image.getexif()
                for key, value in exif.items():
                    label = ExifTags.TAGS.get(key, str(key))
                    exif_values[str(label)] = str(value)
            except Exception:
                exif_values = {}

            metadata = {
                "format": image.format,
                "width": image.width,
                "height": image.height,
                "mode": image.mode,
                "exif": exif_values,
                "needs_vision": True,
            }
            description = (
                f"Image {path.name}: {image.width}x{image.height}, "
                f"format={image.format or path.suffix.lstrip('.').upper()}, mode={image.mode}. "
                "Pixel content is preserved locally; semantic vision analysis is deferred to the model phase."
            )
    except Exception as exc:
        raise ParseError(f"Image parse failed: {exc}") from exc

    return ParsedDocument(
        parser="pillow-metadata",
        parser_version=PARSER_VERSION,
        file_type=path.suffix.lower().lstrip("."),
        title=path.stem,
        text=description,
        units=[ParsedUnit(kind="image", text=description, locator={"image": 1}, metadata=metadata)],
        metadata=metadata,
    )


PARSERS: dict[str, Callable[[Path], ParsedDocument]] = {
    ".txt": parse_text,
    ".md": parse_text,
    ".csv": parse_csv,
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".xlsx": parse_xlsx,
    ".pptx": parse_pptx,
    ".jpg": parse_image,
    ".png": parse_image,
}


def parse_document(path: Path) -> ParsedDocument:
    parser = PARSERS.get(path.suffix.lower())
    if parser is None:
        raise ParseError(f"No Phase 3 parser for {path.suffix.lower() or 'extensionless file'}")
    return parser(path)
