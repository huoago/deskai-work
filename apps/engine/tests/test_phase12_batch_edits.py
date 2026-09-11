from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from sqlalchemy import select

from app.ai.provider import AgentResponse, AgentToolRequest
from app.database.models import SourceEditBatch, SourceFileEdit
from app.security.secrets import SecretStatus


class AgentSecretStore:
    def get_openai_api_key(self) -> str | None:
        return "sk-test-agent-123456789012345678901234"

    def status(self) -> SecretStatus:
        return SecretStatus(
            configured=True,
            source="credential_manager",
            credential_store_available=True,
            writable=True,
        )


class FakeAgentProvider:
    def __init__(self, responses: list[AgentResponse]) -> None:
        self.responses = list(responses)
        self.calls: list[dict] = []

    async def agent_response(self, **kwargs):
        self.calls.append(kwargs)
        if not self.responses:
            raise AssertionError("Fake Agent provider ran out of responses")
        return self.responses.pop(0)


def _tool_step(name: str, arguments: dict, *, call_id: str) -> AgentResponse:
    return AgentResponse(
        text="",
        tool_calls=[AgentToolRequest(call_id=call_id, name=name, arguments=arguments)],
        output_items=[
            {
                "type": "function_call",
                "call_id": call_id,
                "name": name,
                "arguments": json.dumps(arguments, ensure_ascii=False),
            }
        ],
        response_id=f"resp_{call_id}",
        model="gpt-5.6-sol",
        input_tokens=50,
        output_tokens=15,
    )


def _final(text: str) -> AgentResponse:
    return AgentResponse(
        text=text,
        tool_calls=[],
        output_items=[{"type": "message", "role": "assistant"}],
        response_id="resp_final",
        model="gpt-5.6-sol",
        input_tokens=60,
        output_tokens=20,
    )


def _wire_agent(client, provider: FakeAgentProvider) -> None:
    secret_store = AgentSecretStore()
    client.app.state.secret_store = secret_store
    client.app.state.agent_orchestrator.secret_store = secret_store
    client.app.state.agent_orchestrator.provider = provider


def _workspace_with_text_files(client, tmp_path: Path):
    workspace = client.post("/workspaces", json={"name": "Phase 12 batch"}).json()
    root_path = tmp_path / "batch-root"
    root_path.mkdir()
    first = root_path / "first.md"
    second = root_path / "second.txt"
    first.write_text("first old value\n", encoding="utf-8")
    second.write_text("second old value\n", encoding="utf-8")

    root_response = client.post(
        f"/workspaces/{workspace['id']}/roots",
        json={
            "path": str(root_path),
            "read_allowed": True,
            "write_allowed": True,
            "watch_enabled": False,
            "scan_now": True,
        },
    )
    assert root_response.status_code == 201
    files = client.get(f"/files?workspace_id={workspace['id']}").json()
    by_name = {item["filename"]: item for item in files}
    task = client.post(
        "/tasks",
        json={
            "workspace_id": workspace["id"],
            "request": "Update two project source files together",
        },
    ).json()
    return workspace, by_name, task, first, second


def _text_batch(client, workspace, files, task):
    return client.app.state.source_edit_batch_service.propose(
        task_id=task["id"],
        workspace_id=workspace["id"],
        summary="Update coordinated project wording",
        edits=[
            {
                "file_id": files["first.md"]["id"],
                "mode": "text_replace",
                "summary": "Update first file",
                "replacements": [
                    {"find": "old", "replace": "new", "replace_all": False}
                ],
                "cell_edits": [],
            },
            {
                "file_id": files["second.txt"]["id"],
                "mode": "text_replace",
                "summary": "Update second file",
                "replacements": [
                    {"find": "old", "replace": "new", "replace_all": False}
                ],
                "cell_edits": [],
            },
        ],
    )


def test_phase12_batch_confirm_and_rollback_are_all_or_nothing(client, tmp_path):
    workspace, files, task, first, second = _workspace_with_text_files(client, tmp_path)
    batch = _text_batch(client, workspace, files, task)

    assert batch["status"] == "pending"
    assert batch["transactional"] is True
    assert batch["edit_count"] == 2
    assert first.read_text(encoding="utf-8") == "first old value\n"
    assert second.read_text(encoding="utf-8") == "second old value\n"

    detail = client.get(f"/tasks/{task['id']}").json()
    assert len(detail["file_edit_batches"]) == 1
    assert all(item["batch_id"] == batch["id"] for item in detail["file_edits"])

    member_id = batch["edits"][0]["id"]
    member_response = client.post(f"/file-edits/{member_id}/confirm")
    assert member_response.status_code == 409
    assert "transactional batch" in member_response.json()["detail"]

    applied = client.post(f"/file-edit-batches/{batch['id']}/confirm")
    assert applied.status_code == 200
    applied_payload = applied.json()
    assert applied_payload["status"] == "applied"
    assert all(item["status"] == "applied" for item in applied_payload["edits"])
    assert first.read_text(encoding="utf-8") == "first new value\n"
    assert second.read_text(encoding="utf-8") == "second new value\n"

    rolled_back = client.post(f"/file-edit-batches/{batch['id']}/rollback")
    assert rolled_back.status_code == 200
    assert rolled_back.json()["status"] == "rolled_back"
    assert first.read_text(encoding="utf-8") == "first old value\n"
    assert second.read_text(encoding="utf-8") == "second old value\n"


def test_phase12_stale_member_blocks_entire_batch_before_any_write(client, tmp_path):
    workspace, files, task, first, second = _workspace_with_text_files(client, tmp_path)
    batch = _text_batch(client, workspace, files, task)
    second.write_text("external update\n", encoding="utf-8")

    response = client.post(f"/file-edit-batches/{batch['id']}/confirm")
    assert response.status_code == 409
    assert "changed after the batch proposal" in response.json()["detail"]
    assert first.read_text(encoding="utf-8") == "first old value\n"
    assert second.read_text(encoding="utf-8") == "external update\n"
    assert client.get(f"/file-edit-batches/{batch['id']}").json()["status"] == "pending"


def test_phase12_mid_transaction_failure_restores_already_written_files(
    client,
    tmp_path,
    monkeypatch,
):
    workspace, files, task, first, second = _workspace_with_text_files(client, tmp_path)
    batch = _text_batch(client, workspace, files, task)
    edit_service = client.app.state.source_edit_service
    original_atomic_replace = edit_service._atomic_replace
    candidate_writes = 0

    def flaky_replace(source_copy: Path, destination: Path) -> None:
        nonlocal candidate_writes
        if "edit_proposals" in str(source_copy):
            candidate_writes += 1
            if candidate_writes == 2:
                raise OSError("simulated second-file write failure")
        original_atomic_replace(source_copy, destination)

    monkeypatch.setattr(edit_service, "_atomic_replace", flaky_replace)

    response = client.post(f"/file-edit-batches/{batch['id']}/confirm")
    assert response.status_code == 409
    assert "automatically restored" in response.json()["detail"]
    assert first.read_text(encoding="utf-8") == "first old value\n"
    assert second.read_text(encoding="utf-8") == "second old value\n"

    current = client.get(f"/file-edit-batches/{batch['id']}").json()
    assert current["status"] == "pending"
    assert "all already-written files were restored" in current["error_message"]


def test_phase12_office_lock_file_blocks_batch_preflight(client, tmp_path):
    workspace = client.post("/workspaces", json={"name": "Office lock"}).json()
    root_path = tmp_path / "office-root"
    root_path.mkdir()
    docx_path = root_path / "report.docx"
    text_path = root_path / "notes.md"

    document = Document()
    document.add_paragraph("Old report")
    document.save(docx_path)
    text_path.write_text("old notes", encoding="utf-8")

    root_response = client.post(
        f"/workspaces/{workspace['id']}/roots",
        json={
            "path": str(root_path),
            "read_allowed": True,
            "write_allowed": True,
            "watch_enabled": False,
            "scan_now": True,
        },
    )
    assert root_response.status_code == 201
    files = client.get(f"/files?workspace_id={workspace['id']}").json()
    by_name = {item["filename"]: item for item in files}
    task = client.post(
        "/tasks",
        json={"workspace_id": workspace["id"], "request": "Update report and notes together"},
    ).json()

    batch = client.app.state.source_edit_batch_service.propose(
        task_id=task["id"],
        workspace_id=workspace["id"],
        summary="Coordinated Office update",
        edits=[
            {
                "file_id": by_name["report.docx"]["id"],
                "mode": "docx_replace",
                "summary": "Update report",
                "replacements": [
                    {"find": "Old", "replace": "New", "replace_all": False}
                ],
                "cell_edits": [],
            },
            {
                "file_id": by_name["notes.md"]["id"],
                "mode": "text_replace",
                "summary": "Update notes",
                "replacements": [
                    {"find": "old", "replace": "new", "replace_all": False}
                ],
                "cell_edits": [],
            },
        ],
    )

    lock_path = root_path / "~$report.docx"
    lock_path.write_text("office lock", encoding="utf-8")
    response = client.post(f"/file-edit-batches/{batch['id']}/confirm")
    assert response.status_code == 409
    assert "open in Microsoft Office" in response.json()["detail"]
    assert Document(docx_path).paragraphs[0].text == "Old report"
    assert text_path.read_text(encoding="utf-8") == "old notes"


def test_phase12_startup_recovery_restores_interrupted_apply(client, tmp_path):
    workspace, files, task, first, second = _workspace_with_text_files(client, tmp_path)
    batch = _text_batch(client, workspace, files, task)
    batch_service = client.app.state.source_edit_batch_service
    batch_record = batch_service._batch_record(batch["id"], expected_status="pending")
    snapshots = batch_service._snapshots(batch_record, expected_edit_status="pending")
    batch_service._preflight_apply(snapshots)
    batch_service._create_backups(snapshots)

    now = datetime.now(timezone.utc)
    with client.app.state.database.session() as session:
        current = session.get(SourceEditBatch, batch["id"])
        assert current is not None
        current.status = "applying"
        current.confirmed_at = now
        members = list(
            session.scalars(
                select(SourceFileEdit).where(SourceFileEdit.batch_id == batch["id"])
            ).all()
        )
        for edit in members:
            edit.status = "applying"
            edit.confirmed_at = now
            snapshot = next(item for item in snapshots if item["edit_id"] == edit.id)
            edit.backup_path = str(snapshot["backup"])

    client.app.state.source_edit_service._atomic_replace(
        snapshots[0]["candidate"],
        snapshots[0]["source"],
    )
    assert first.read_text(encoding="utf-8") == "first new value\n"
    assert second.read_text(encoding="utf-8") == "second old value\n"

    recovered = batch_service.recover_incomplete_batches()
    assert recovered == {"recovered": 1, "blocked": 0}
    assert first.read_text(encoding="utf-8") == "first old value\n"
    assert second.read_text(encoding="utf-8") == "second old value\n"
    current = client.get(f"/file-edit-batches/{batch['id']}").json()
    assert current["status"] == "pending"
    assert "recovered an interrupted batch apply" in current["error_message"]


def test_phase12_agent_can_only_stage_transactional_batch(client, tmp_path):
    workspace, files, task, first, second = _workspace_with_text_files(client, tmp_path)
    arguments = {
        "summary": "Update both policy files",
        "edits": [
            {
                "file_id": files["first.md"]["id"],
                "mode": "text_replace",
                "summary": "Update first",
                "replacements": [
                    {"find": "old", "replace": "new", "replace_all": False}
                ],
                "cell_edits": [],
            },
            {
                "file_id": files["second.txt"]["id"],
                "mode": "text_replace",
                "summary": "Update second",
                "replacements": [
                    {"find": "old", "replace": "new", "replace_all": False}
                ],
                "cell_edits": [],
            },
        ],
    }
    provider = FakeAgentProvider(
        [
            _tool_step(
                "propose_source_file_edit_batch",
                arguments,
                call_id="batch_1",
            ),
            _final("批量编辑事务已生成，正在等待你一次确认整个变更集。"),
        ]
    )
    _wire_agent(client, provider)

    assert client.post("/agent/process", params={"limit": 1}).json()["processed"] == 1
    detail = client.get(f"/tasks/{task['id']}").json()
    assert detail["status"] == "completed"
    assert first.read_text(encoding="utf-8") == "first old value\n"
    assert second.read_text(encoding="utf-8") == "second old value\n"
    assert len(detail["file_edit_batches"]) == 1
    batch = detail["file_edit_batches"][0]
    assert batch["status"] == "pending"
    assert batch["transactional"] is True
    assert batch["edit_count"] == 2
    assert all(item["batch_id"] == batch["id"] for item in batch["edits"])
    assert [call["tool_name"] for call in detail["tool_calls"]] == [
        "propose_source_file_edit_batch"
    ]
    assert detail["tool_calls"][0]["risk_level"] == 3

    second_items = provider.calls[1]["input_items"]
    output = next(
        item
        for item in second_items
        if item.get("type") == "function_call_output" and item.get("call_id") == "batch_1"
    )
    payload = json.loads(output["output"])
    assert payload["result"]["requires_user_confirmation"] is True
    assert payload["result"]["transactional"] is True


def test_phase12_rejects_duplicate_file_in_same_batch(client, tmp_path):
    workspace, files, task, _first, _second = _workspace_with_text_files(client, tmp_path)
    duplicate = {
        "file_id": files["first.md"]["id"],
        "mode": "text_replace",
        "summary": "Duplicate",
        "replacements": [{"find": "old", "replace": "new", "replace_all": False}],
        "cell_edits": [],
    }

    try:
        client.app.state.source_edit_batch_service.propose(
            task_id=task["id"],
            workspace_id=workspace["id"],
            summary="Invalid duplicate batch",
            edits=[duplicate, duplicate],
        )
    except ValueError as exc:
        assert "same file more than once" in str(exc)
    else:
        raise AssertionError("duplicate file batch should have been rejected")
