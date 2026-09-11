from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from app.ai.provider import AgentResponse, AgentToolRequest
from app.security.secrets import SecretStatus


class AgentSecretStore:
    def get_openai_api_key(self) -> str | None:
        return "sk-test-agent-123456789012345678901234"

    def status(self) -> SecretStatus:
        return SecretStatus(
            configured=True,
            source="credential_manager",
            credential_store_available=True,
            writable=True,
        )


class FakeAgentProvider:
    def __init__(self, responses: list[AgentResponse]) -> None:
        self.responses = list(responses)
        self.calls: list[dict] = []

    async def agent_response(self, **kwargs):
        self.calls.append(kwargs)
        if not self.responses:
            raise AssertionError("Fake Agent provider ran out of responses")
        return self.responses.pop(0)


def _wire_provider(client, provider: FakeAgentProvider) -> None:
    secret_store = AgentSecretStore()
    client.app.state.secret_store = secret_store
    client.app.state.agent_orchestrator.secret_store = secret_store
    client.app.state.agent_orchestrator.provider = provider


def _tool_step(name: str, arguments: dict, *, call_id: str) -> AgentResponse:
    return AgentResponse(
        text="",
        tool_calls=[AgentToolRequest(call_id=call_id, name=name, arguments=arguments)],
        output_items=[
            {
                "type": "function_call",
                "call_id": call_id,
                "name": name,
                "arguments": json.dumps(arguments, ensure_ascii=False),
            }
        ],
        response_id=f"resp_{call_id}",
        model="gpt-5.6-sol",
        input_tokens=40,
        output_tokens=10,
    )


def _final(text: str) -> AgentResponse:
    return AgentResponse(
        text=text,
        tool_calls=[],
        output_items=[{"type": "message", "role": "assistant"}],
        response_id="resp_final",
        model="gpt-5.6-sol",
        input_tokens=50,
        output_tokens=20,
    )


def test_phase8_agent_creates_real_docx_in_private_task_sandbox(client):
    workspace = client.post("/workspaces", json={"name": "Artifact Project"}).json()
    arguments = {
        "filename": "../../Desktop/项目总结.docx",
        "title": "324二级给水工作总结",
        "sections": [
            {
                "heading": "一、结论",
                "body": "水表统计总数为4447个。\n后续需继续核对正式批复。",
            },
            {
                "heading": "二、注意事项",
                "body": "本文件由DeskAI生成。",
            },
        ],
    }
    provider = FakeAgentProvider(
        [
            _tool_step("create_word_document", arguments, call_id="call_docx"),
            _final("已生成Word工作成果。"),
        ]
    )
    _wire_provider(client, provider)

    task = client.post(
        "/tasks",
        json={
            "workspace_id": workspace["id"],
            "request": "把结论整理成Word文件。",
        },
    ).json()
    assert client.post("/agent/process", params={"limit": 1}).json()["processed"] == 1

    detail = client.get(f"/tasks/{task['id']}").json()
    assert detail["status"] == "completed"
    assert len(detail["artifacts"]) == 1
    artifact = detail["artifacts"][0]
    assert artifact["kind"] == "docx"
    assert artifact["filename"] == "项目总结.docx"

    path = Path(artifact["path"]).resolve()
    generated_root = client.app.state.artifact_service.root.resolve()
    task_root = (generated_root / task["id"]).resolve()
    assert path.is_relative_to(task_root)
    assert "Desktop" not in path.parts
    assert path.exists()

    document = Document(path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "324二级给水工作总结" in text
    assert "水表统计总数为4447个" in text
    assert "本文件由DeskAI生成" in text

    model_input = json.dumps(provider.calls[1]["input_items"], ensure_ascii=False)
    assert artifact["filename"] in model_input
    assert str(path) not in model_input
    assert str(generated_root) not in model_input

    listed = client.get(
        "/artifacts",
        params={"workspace_id": workspace["id"], "task_id": task["id"]},
    )
    assert listed.status_code == 200
    assert listed.json()[0]["id"] == artifact["id"]


def test_phase8_duplicate_artifact_names_never_overwrite(client):
    workspace = client.post("/workspaces", json={"name": "No Overwrite"}).json()
    first = {
        "filename": "报告.docx",
        "title": "第一版",
        "sections": [{"heading": "内容", "body": "first"}],
    }
    second = {
        "filename": "报告.docx",
        "title": "第二版",
        "sections": [{"heading": "内容", "body": "second"}],
    }
    provider = FakeAgentProvider(
        [
            _tool_step("create_word_document", first, call_id="call_1"),
            _tool_step("create_word_document", second, call_id="call_2"),
            _final("已生成两个独立版本。"),
        ]
    )
    _wire_provider(client, provider)

    task = client.post(
        "/tasks",
        json={"workspace_id": workspace["id"], "request": "分别生成两个Word版本。"},
    ).json()
    assert client.post("/agent/process", params={"limit": 1}).json()["processed"] == 1

    detail = client.get(f"/tasks/{task['id']}").json()
    assert detail["status"] == "completed"
    assert len(detail["artifacts"]) == 2
    names = {item["filename"] for item in detail["artifacts"]}
    assert "报告.docx" in names
    assert len(names) == 2

    contents = []
    for artifact in detail["artifacts"]:
        path = Path(artifact["path"])
        assert path.exists()
        document = Document(path)
        contents.append("\n".join(p.text for p in document.paragraphs))
    assert any("第一版" in text and "first" in text for text in contents)
    assert any("第二版" in text and "second" in text for text in contents)


def test_phase8_agent_creates_real_xlsx_and_neutralizes_formula_injection(client):
    workspace = client.post("/workspaces", json={"name": "Spreadsheet Safety"}).json()
    arguments = {
        "filename": "../危险统计.xlsx",
        "sheets": [
            {
                "name": "统计/表",
                "headers": ["项目", "值", "备注"],
                "rows": [
                    ["总数", "4447", "=2+2"],
                    ["异常", "4", "@cmd"],
                    ["普通", "154", "正常文本"],
                ],
            }
        ],
    }
    provider = FakeAgentProvider(
        [
            _tool_step("create_spreadsheet", arguments, call_id="call_xlsx"),
            _final("已生成Excel统计表。"),
        ]
    )
    _wire_provider(client, provider)

    task = client.post(
        "/tasks",
        json={"workspace_id": workspace["id"], "request": "生成Excel统计表。"},
    ).json()
    assert client.post("/agent/process", params={"limit": 1}).json()["processed"] == 1

    detail = client.get(f"/tasks/{task['id']}").json()
    assert detail["status"] == "completed"
    artifact = detail["artifacts"][0]
    assert artifact["kind"] == "xlsx"
    assert artifact["filename"] == "危险统计.xlsx"

    path = Path(artifact["path"]).resolve()
    task_root = (client.app.state.artifact_service.root / task["id"]).resolve()
    assert path.is_relative_to(task_root)
    workbook = load_workbook(path, data_only=False)
    sheet = workbook[workbook.sheetnames[0]]
    assert sheet.title == "统计_表"
    assert sheet["A2"].value == "总数"
    assert sheet["B2"].value == "4447"
    assert sheet["C2"].value == "'=2+2"
    assert sheet["C2"].data_type == "s"
    assert sheet["C3"].value == "'@cmd"
    assert sheet["C3"].data_type == "s"
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref is not None


def test_phase8_artifact_listing_is_workspace_scoped(client):
    first = client.post("/workspaces", json={"name": "Artifacts A"}).json()
    second = client.post("/workspaces", json={"name": "Artifacts B"}).json()

    def create_for(workspace_id: str, filename: str):
        provider = FakeAgentProvider(
            [
                _tool_step(
                    "create_word_document",
                    {
                        "filename": filename,
                        "title": filename,
                        "sections": [{"heading": "内容", "body": filename}],
                    },
                    call_id=f"call_{filename}",
                ),
                _final("完成"),
            ]
        )
        _wire_provider(client, provider)
        task = client.post(
            "/tasks",
            json={"workspace_id": workspace_id, "request": f"生成{filename}"},
        ).json()
        assert client.post("/agent/process", params={"limit": 1}).json()["processed"] == 1
        return task

    first_task = create_for(first["id"], "A.docx")
    create_for(second["id"], "B.docx")

    first_items = client.get("/artifacts", params={"workspace_id": first["id"]}).json()
    assert len(first_items) == 1
    assert first_items[0]["filename"] == "A.docx"
    assert first_items[0]["task_id"] == first_task["id"]
