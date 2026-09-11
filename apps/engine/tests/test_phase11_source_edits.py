from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

from app.ai.provider import AgentResponse, AgentToolRequest
from app.security.secrets import SecretStatus


class AgentSecretStore:
    def get_openai_api_key(self) -> str | None:
        return "sk-test-agent-123456789012345678901234"

    def status(self) -> SecretStatus:
        return SecretStatus(
            configured=True,
            source="credential_manager",
            credential_store_available=True,
            writable=True,
        )


class FakeAgentProvider:
    def __init__(self, responses: list[AgentResponse]) -> None:
        self.responses = list(responses)
        self.calls: list[dict] = []

    async def agent_response(self, **kwargs):
        self.calls.append(kwargs)
        if not self.responses:
            raise AssertionError("Fake Agent provider ran out of responses")
        return self.responses.pop(0)


def _tool_step(name: str, arguments: dict, *, call_id: str) -> AgentResponse:
    return AgentResponse(
        text="",
        tool_calls=[AgentToolRequest(call_id=call_id, name=name, arguments=arguments)],
        output_items=[
            {
                "type": "function_call",
                "call_id": call_id,
                "name": name,
                "arguments": json.dumps(arguments, ensure_ascii=False),
            }
        ],
        response_id=f"resp_{call_id}",
        model="gpt-5.6-sol",
        input_tokens=40,
        output_tokens=10,
    )


def _final(text: str) -> AgentResponse:
    return AgentResponse(
        text=text,
        tool_calls=[],
        output_items=[{"type": "message", "role": "assistant"}],
        response_id="resp_final",
        model="gpt-5.6-sol",
        input_tokens=50,
        output_tokens=20,
    )


def _wire_agent(client, provider: FakeAgentProvider) -> None:
    secret_store = AgentSecretStore()
    client.app.state.secret_store = secret_store
    client.app.state.agent_orchestrator.secret_store = secret_store
    client.app.state.agent_orchestrator.provider = provider


def _workspace_with_file(
    client,
    tmp_path: Path,
    *,
    filename: str,
    create_file,
    write_allowed: bool = False,
):
    workspace = client.post("/workspaces", json={"name": f"Edit {filename}"}).json()
    root_path = tmp_path / f"root-{filename.replace('.', '-')}"
    root_path.mkdir()
    source = root_path / filename
    create_file(source)
    response = client.post(
        f"/workspaces/{workspace['id']}/roots",
        json={
            "path": str(root_path),
            "read_allowed": True,
            "write_allowed": write_allowed,
            "watch_enabled": False,
            "scan_now": True,
        },
    )
    assert response.status_code == 201
    root = response.json()["root"]
    files = client.get(f"/files?workspace_id={workspace['id']}").json()
    file = next(item for item in files if item["filename"] == filename)
    task = client.post(
        "/tasks",
        json={
            "workspace_id": workspace["id"],
            "request": f"Modify {filename}",
        },
    ).json()
    return workspace, root, file, task, source


def test_phase11_text_edit_requires_write_permission_and_human_confirmation(client, tmp_path):
    workspace, root, file, task, source = _workspace_with_file(
        client,
        tmp_path,
        filename="notes.md",
        create_file=lambda path: path.write_text("alpha old beta\n", encoding="utf-8"),
        write_allowed=False,
    )

    with pytest.raises(ValueError, match="write_allowed"):
        client.app.state.source_edit_service.propose(
            task_id=task["id"],
            workspace_id=workspace["id"],
            file_id=file["id"],
            mode="text_replace",
            summary="Replace old with new",
            replacements=[{"find": "old", "replace": "new", "replace_all": True}],
            cell_edits=[],
        )

    update = client.patch(
        f"/workspaces/{workspace['id']}/roots/{root['id']}",
        json={"write_allowed": True},
    )
    assert update.status_code == 200

    proposal = client.app.state.source_edit_service.propose(
        task_id=task["id"],
        workspace_id=workspace["id"],
        file_id=file["id"],
        mode="text_replace",
        summary="Replace old with new",
        replacements=[{"find": "old", "replace": "new", "replace_all": True}],
        cell_edits=[],
    )
    assert proposal["status"] == "pending"
    assert proposal["requires_user_confirmation"] is True
    assert "-alphaoldbeta" in proposal["diff_preview"].replace(" ", "")
    assert source.read_text(encoding="utf-8") == "alpha old beta\n"

    detail = client.get(f"/tasks/{task['id']}").json()
    assert detail["file_edits"][0]["id"] == proposal["id"]

    applied_response = client.post(f"/file-edits/{proposal['id']}/confirm")
    assert applied_response.status_code == 200
    applied = applied_response.json()
    assert applied["status"] == "applied"
    assert applied["backup_created"] is True
    assert source.read_text(encoding="utf-8") == "alpha new beta\n"

    rollback_response = client.post(f"/file-edits/{proposal['id']}/rollback")
    assert rollback_response.status_code == 200
    assert rollback_response.json()["status"] == "rolled_back"
    assert source.read_text(encoding="utf-8") == "alpha old beta\n"


def test_phase11_confirmation_blocks_stale_source_overwrite(client, tmp_path):
    workspace, root, file, task, source = _workspace_with_file(
        client,
        tmp_path,
        filename="status.txt",
        create_file=lambda path: path.write_text("version one", encoding="utf-8"),
        write_allowed=True,
    )
    assert root["write_allowed"] is True

    proposal = client.app.state.source_edit_service.propose(
        task_id=task["id"],
        workspace_id=workspace["id"],
        file_id=file["id"],
        mode="text_replace",
        summary="Update version",
        replacements=[{"find": "one", "replace": "two", "replace_all": False}],
        cell_edits=[],
    )
    source.write_text("external change", encoding="utf-8")

    response = client.post(f"/file-edits/{proposal['id']}/confirm")
    assert response.status_code == 409
    assert "changed after the proposal" in response.json()["detail"]
    assert source.read_text(encoding="utf-8") == "external change"
    assert client.get(f"/file-edits/{proposal['id']}").json()["status"] == "pending"


def test_phase11_docx_edit_is_staged_before_apply(client, tmp_path):
    def create_docx(path: Path) -> None:
        document = Document()
        document.add_paragraph("Before replacement")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Before table"
        document.save(path)

    workspace, _root, file, task, source = _workspace_with_file(
        client,
        tmp_path,
        filename="report.docx",
        create_file=create_docx,
        write_allowed=True,
    )
    proposal = client.app.state.source_edit_service.propose(
        task_id=task["id"],
        workspace_id=workspace["id"],
        file_id=file["id"],
        mode="docx_replace",
        summary="Replace wording",
        replacements=[{"find": "Before", "replace": "After", "replace_all": True}],
        cell_edits=[],
    )

    original = Document(source)
    assert original.paragraphs[0].text == "Before replacement"
    assert proposal["status"] == "pending"

    assert client.post(f"/file-edits/{proposal['id']}/confirm").status_code == 200
    changed = Document(source)
    assert changed.paragraphs[0].text == "After replacement"
    assert changed.tables[0].cell(0, 0).text == "After table"


def test_phase11_xlsx_edit_neutralizes_formula_like_text(client, tmp_path):
    def create_xlsx(path: Path) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data"
        sheet["A1"] = "old"
        sheet["B1"] = 10
        workbook.save(path)

    workspace, _root, file, task, source = _workspace_with_file(
        client,
        tmp_path,
        filename="metrics.xlsx",
        create_file=create_xlsx,
        write_allowed=True,
    )
    proposal = client.app.state.source_edit_service.propose(
        task_id=task["id"],
        workspace_id=workspace["id"],
        file_id=file["id"],
        mode="xlsx_cells",
        summary="Update cells",
        replacements=[],
        cell_edits=[
            {"sheet": "Data", "cell": "A1", "value_type": "text", "value": "=1+1"},
            {"sheet": "Data", "cell": "B1", "value_type": "number", "value": "25"},
        ],
    )
    assert proposal["status"] == "pending"
    before = load_workbook(source, data_only=False)
    assert before["Data"]["A1"].value == "old"

    assert client.post(f"/file-edits/{proposal['id']}/confirm").status_code == 200
    after = load_workbook(source, data_only=False)
    assert after["Data"]["A1"].value == "'=1+1"
    assert after["Data"]["B1"].value == 25


def test_phase11_agent_can_only_create_pending_edit_proposal(client, tmp_path):
    workspace, _root, file, task, source = _workspace_with_file(
        client,
        tmp_path,
        filename="agent.md",
        create_file=lambda path: path.write_text("old policy", encoding="utf-8"),
        write_allowed=True,
    )
    provider = FakeAgentProvider(
        [
            _tool_step(
                "propose_source_file_edit",
                {
                    "file_id": file["id"],
                    "mode": "text_replace",
                    "summary": "Update policy wording",
                    "replacements": [
                        {"find": "old", "replace": "new", "replace_all": False}
                    ],
                    "cell_edits": [],
                },
                call_id="edit_1",
            ),
            _final("编辑提案已经生成，正在等待你在任务详情中确认应用。"),
        ]
    )
    _wire_agent(client, provider)

    assert client.post("/agent/process", params={"limit": 1}).json()["processed"] == 1
    detail = client.get(f"/tasks/{task['id']}").json()
    assert detail["status"] == "completed"
    assert source.read_text(encoding="utf-8") == "old policy"
    assert len(detail["file_edits"]) == 1
    assert detail["file_edits"][0]["status"] == "pending"
    assert [call["tool_name"] for call in detail["tool_calls"]] == [
        "propose_source_file_edit"
    ]
    assert detail["tool_calls"][0]["risk_level"] == 3

    second_items = provider.calls[1]["input_items"]
    output = next(
        item
        for item in second_items
        if item.get("type") == "function_call_output" and item.get("call_id") == "edit_1"
    )
    payload = json.loads(output["output"])
    assert payload["result"]["requires_user_confirmation"] is True
    assert payload["result"]["status"] == "pending"
