from __future__ import annotations

from pathlib import Path

import fitz
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from sqlalchemy import select

from app.database.models import Chunk, File, FileVersion


def _workspace(client) -> str:
    response = client.post("/workspaces", json={"name": "Parser Test"})
    assert response.status_code == 201
    return response.json()["id"]


def _authorize_and_parse(client, workspace_id: str, source: Path) -> list[dict]:
    added = client.post(
        f"/workspaces/{workspace_id}/roots",
        json={"path": str(source), "scan_now": True, "watch_enabled": False},
    )
    assert added.status_code == 201
    queued = added.json()["scan"]["queued"]
    assert queued > 0
    processed = client.post("/parser/process", params={"limit": 50})
    assert processed.status_code == 200
    assert processed.json()["processed"] == queued
    return client.get("/files", params={"workspace_id": workspace_id}).json()


def _preview(client, file: dict) -> dict:
    response = client.get(f"/files/{file['id']}/parsed")
    assert response.status_code == 200
    payload = response.json()
    assert payload["parser_version"] == "phase3.1"
    assert payload["sha256"] == file["sha256"]
    return payload


def test_phase3_parses_text_markdown_and_csv(client, tmp_path: Path):
    workspace_id = _workspace(client)
    source = tmp_path / "text"
    source.mkdir()
    (source / "note.txt").write_text("RRP-04 pressure test\n\nAccepted.", encoding="utf-8")
    (source / "manual.md").write_text("# Section 313\nPipeline length 284 m.", encoding="utf-8")
    (source / "meters.csv").write_text("sector,count\n324,4447\n", encoding="utf-8")

    files = _authorize_and_parse(client, workspace_id, source)
    assert {item["status"] for item in files} == {"parsed"}

    by_name = {item["filename"]: item for item in files}
    assert "RRP-04 pressure test" in _preview(client, by_name["note.txt"])["text"]
    assert "Pipeline length 284 m." in _preview(client, by_name["manual.md"])["text"]
    csv_preview = _preview(client, by_name["meters.csv"])
    assert "324\t4447" in csv_preview["text"]
    assert csv_preview["metadata"]["row_count"] == 2


def test_phase3_parses_pdf_pages(client, tmp_path: Path):
    workspace_id = _workspace(client)
    source = tmp_path / "pdf"
    source.mkdir()
    path = source / "handover.pdf"

    document = fitz.open()
    page1 = document.new_page()
    page1.insert_text((72, 72), "Sector 324 handover")
    page2 = document.new_page()
    page2.insert_text((72, 72), "Pressure test approved")
    document.save(path)
    document.close()

    files = _authorize_and_parse(client, workspace_id, source)
    payload = _preview(client, files[0])
    assert payload["parser"] == "pymupdf"
    assert "Sector 324 handover" in payload["text"]
    assert "Pressure test approved" in payload["text"]
    assert [unit["locator"]["page"] for unit in payload["units"]] == [1, 2]


def test_phase3_parses_docx_paragraphs_and_tables(client, tmp_path: Path):
    workspace_id = _workspace(client)
    source = tmp_path / "docx"
    source.mkdir()
    path = source / "report.docx"

    document = Document()
    document.add_heading("313 Construction Progress", level=1)
    document.add_paragraph("Installed pipeline: 40 m")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Manholes"
    table.cell(1, 1).text = "5"
    document.save(path)

    files = _authorize_and_parse(client, workspace_id, source)
    payload = _preview(client, files[0])
    assert payload["parser"] == "python-docx"
    assert "313 Construction Progress" in payload["text"]
    assert "Manholes\t5" in payload["text"]
    assert any(unit["kind"] == "table_row" for unit in payload["units"])


def test_phase3_parses_xlsx_sheets_and_formulas(client, tmp_path: Path):
    workspace_id = _workspace(client)
    source = tmp_path / "xlsx"
    source.mkdir()
    path = source / "meters.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "324"
    sheet.append(["Category", "Count"])
    sheet.append([">5 years", 1802])
    sheet.append(["Total", "=SUM(B2:B2)"])
    second = workbook.create_sheet("313")
    second.append(["Length", 284])
    workbook.save(path)
    workbook.close()

    files = _authorize_and_parse(client, workspace_id, source)
    payload = _preview(client, files[0])
    assert payload["parser"] == "openpyxl"
    assert payload["metadata"]["sheet_count"] == 2
    assert ">5 years\t1802" in payload["text"]
    assert "=SUM(B2:B2)" in payload["text"]
    assert any(unit["locator"]["sheet"] == "313" for unit in payload["units"])


def test_phase3_parses_pptx_slides(client, tmp_path: Path):
    workspace_id = _workspace(client)
    source = tmp_path / "pptx"
    source.mkdir()
    path = source / "briefing.pptx"

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.text = "CD-03 hot tapping"
    slide2 = presentation.slides.add_slide(presentation.slide_layouts[5])
    box2 = slide2.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box2.text = "DN1500 under pressure"
    presentation.save(path)

    files = _authorize_and_parse(client, workspace_id, source)
    payload = _preview(client, files[0])
    assert payload["parser"] == "python-pptx"
    assert payload["metadata"]["slide_count"] == 2
    assert "CD-03 hot tapping" in payload["text"]
    assert "DN1500 under pressure" in payload["text"]
    assert [unit["locator"]["slide"] for unit in payload["units"]] == [1, 2]


def test_phase3_parses_image_metadata_without_fake_ocr(client, tmp_path: Path):
    workspace_id = _workspace(client)
    source = tmp_path / "image"
    source.mkdir()
    path = source / "site.png"
    Image.new("RGB", (320, 240)).save(path)

    files = _authorize_and_parse(client, workspace_id, source)
    payload = _preview(client, files[0])
    assert payload["parser"] == "pillow-metadata"
    assert payload["metadata"]["width"] == 320
    assert payload["metadata"]["height"] == 240
    assert payload["metadata"]["needs_vision"] is True
    assert "semantic vision analysis is deferred" in payload["text"]


def test_phase3_creates_file_version_but_not_chunks(client, tmp_path: Path):
    workspace_id = _workspace(client)
    source = tmp_path / "version"
    source.mkdir()
    path = source / "versioned.txt"
    path.write_text("version one", encoding="utf-8")

    files = _authorize_and_parse(client, workspace_id, source)
    first_file = files[0]
    assert first_file["current_version_id"]
    assert first_file["parser_version"] == "phase3.1"

    with client.app.state.database.session() as session:
        versions = session.scalars(select(FileVersion)).all()
        chunks = session.scalars(select(Chunk)).all()
        assert len(versions) == 1
        assert versions[0].active is True
        assert chunks == []

    path.write_text("version two with changed length", encoding="utf-8")
    scan = client.post(f"/workspaces/{workspace_id}/scan").json()
    assert scan["summary"]["queued"] == 1
    client.post("/parser/process", params={"limit": 10})

    with client.app.state.database.session() as session:
        versions = session.scalars(select(FileVersion).order_by(FileVersion.id)).all()
        current_file = session.get(File, first_file["id"])
        assert len(versions) == 2
        assert sum(1 for version in versions if version.active) == 1
        assert current_file is not None
        assert current_file.status == "parsed"
        assert current_file.current_version_id in {version.id for version in versions if version.active}


def test_parser_status_reports_completed_work(client, tmp_path: Path):
    workspace_id = _workspace(client)
    source = tmp_path / "status"
    source.mkdir()
    (source / "one.txt").write_text("hello", encoding="utf-8")

    _authorize_and_parse(client, workspace_id, source)
    status = client.get("/parser/status").json()
    assert status["processed"] == 1
    assert status["failed"] == 0
    assert status["parser_version"] == "phase3.1"
